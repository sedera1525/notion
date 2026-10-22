"""Fonctions d'aide partagées par les générateurs de rapports fictifs.

Utilisé par generer_rapport_demo.py et generer_rapport_cas_c.py.

Le préfixe « _ » signale un module interne aux scripts : il n'est pas destiné
à être exécuté directement.
"""

from docx.shared import Pt


def preparer(doc):
    """Applique la police commune aux deux rapports."""
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    return doc


def titre(doc, texte, niveau):
    doc.add_heading(texte, level=niveau)


def para(doc, texte, gras=False, italique=False):
    p = doc.add_paragraph()
    s = p.add_run(texte)
    s.bold, s.italic = gras, italique
    return p


def tableau(doc, entetes, lignes):
    t = doc.add_table(rows=1, cols=len(entetes))
    t.style = "Light Grid Accent 1"
    for cell, texte in zip(t.rows[0].cells, entetes):
        cell.paragraphs[0].add_run(texte).bold = True
    for ligne in lignes:
        for cell, texte in zip(t.add_row().cells, ligne):
            cell.text = texte
    doc.add_paragraph()
    return t
