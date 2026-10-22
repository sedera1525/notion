#!/usr/bin/env python3
"""Génère le rapport d'activité fictif utilisé pour la démonstration.

Toutes les données sont inventées. Aucune entreprise, aucun montant et
aucune personne ne correspond à la réalité.

⚠️ Le fichier produit n'est pas stable octet par octet : python-docx horodate
les entrées de l'archive zip, donc deux générations donnent des empreintes
SHA-256 différentes. Comme lancer_synthese.sh suit les rapports par empreinte,
régénérer ce fichier le fait passer pour un rapport corrigé, et le prochain
traitement ajoutera une section « mise à jour » dans Notion.

Usage : .venv/bin/python scripts/generer_rapport_demo.py
Sortie : input/rapport-t3-demo.docx
"""

from pathlib import Path

from _docx_commun import para, preparer, tableau, titre
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

SORTIE = Path(__file__).resolve().parent.parent / "input" / "rapport-t3-demo.docx"


def construire():
    doc = preparer(Document())

    # ---------- Page de garde ----------
    t = doc.add_heading("Rapport d'activité trimestriel", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st = para(doc, "Vanille Horizon — Structuration de la filière vanille durable", gras=True)
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    tableau(doc, ["Rubrique", "Information"], [
        ["Client", "Vanille Horizon"],
        ["Pays", "Madagascar"],
        ["Période couverte", "1er juillet 2026 – 30 septembre 2026 (T3 2026)"],
        ["Auteur du rapport", "R. Andriamalala, chargée d'études"],
        ["Date de remise", "8 octobre 2026"],
        ["Référence", "VH-2026-T3-001"],
    ])
    para(doc, "Document de démonstration. Données entièrement fictives.", italique=True)
    doc.paragraphs[-1].runs[0].add_break(WD_BREAK.PAGE)

    # ---------- 1. Contexte ----------
    titre(doc, "1. Contexte et objectifs", 1)
    para(doc,
         "Vanille Horizon accompagne depuis 2024 la structuration d'une filière vanille durable "
         "dans la région Nord-Est de Madagascar. La mission vise trois objectifs : sécuriser le "
         "revenu des producteurs partenaires, améliorer la traçabilité des lots exportés, et "
         "obtenir la certification biologique sur la majorité des surfaces cultivées.")
    para(doc,
         "Le présent rapport couvre le troisième trimestre 2026. Il fait suite au rapport du "
         "deuxième trimestre, remis le 10 juillet 2026. Il s'adresse au comité de pilotage, qui "
         "se réunit le 14 novembre 2026.")
    para(doc,
         "Le trimestre a été marqué par l'extension du réseau de collecte à deux nouvelles "
         "communes et par le lancement du programme de formation des producteurs.")

    # ---------- 2. Chiffres clés ----------
    titre(doc, "2. Chiffres clés du trimestre", 1)
    para(doc, "Les indicateurs ci-dessous sont comparés au deuxième trimestre 2026 lorsque "
              "la donnée antérieure est disponible.")
    tableau(doc, ["Indicateur", "T3 2026", "T2 2026", "Évolution"], [
        ["Chiffre d'affaires", "412 000 €", "358 000 €", "+15,1 %"],
        ["Producteurs partenaires", "247", "198", "+49 producteurs"],
        ["Volume de vanille verte collecté", "18,4 tonnes", "14,2 tonnes", "+29,6 %"],
        ["Surfaces certifiées biologiques", "62 %", "54 %", "+8 points"],
        ["Participants aux formations", "86", "0", "programme lancé au T3"],
        ["Délai moyen de paiement des producteurs", "21 jours", "26 jours", "−5 jours"],
        ["Taux de perte après récolte", "7,3 %", "—", "—"],
    ])
    para(doc,
         "Le taux de perte après récolte est mesuré pour la première fois ce trimestre. "
         "Aucune comparaison n'est possible avec les périodes précédentes.")
    para(doc,
         "L'objectif interne de délai de paiement reste fixé à 15 jours. Il n'est pas atteint.")

    # ---------- 3. Réalisations ----------
    titre(doc, "3. Réalisations de la période", 1)
    titre(doc, "3.1 Extension du réseau de collecte", 2)
    para(doc,
         "Deux centres de collecte ont été ouverts en juillet 2026. Ils couvrent désormais "
         "deux communes supplémentaires, portant le total à sept communes. Le réseau a permis "
         "d'enregistrer 49 nouveaux producteurs partenaires sur le trimestre.")
    para(doc,
         "Chaque centre est équipé d'une balance certifiée et d'un registre numérique de "
         "pesée. La saisie est effectuée par un agent formé sur place.")

    titre(doc, "3.2 Traçabilité des lots", 2)
    para(doc,
         "Le système d'identification par lot est opérationnel depuis le 1er août 2026. "
         "Sur les 18,4 tonnes collectées, 16,1 tonnes sont associées à un identifiant de lot "
         "complet, soit 87,5 % du volume. Le reste correspond aux collectes antérieures à la "
         "mise en service.")

    titre(doc, "3.3 Certification biologique", 2)
    para(doc,
         "Les audits de la deuxième vague ont été menés en septembre 2026 sur 41 exploitations. "
         "38 ont été validées. Les trois exploitations restantes doivent corriger des écarts "
         "mineurs de tenue de registre avant un nouveau passage.")

    # ---------- 4. Formation ----------
    titre(doc, "4. Programme de formation", 1)
    para(doc,
         "Trois sessions de formation ont été organisées au cours du trimestre, pour un total "
         "de 86 participants.")
    tableau(doc, ["Session", "Dates", "Thème", "Participants"], [
        ["Session 1", "22–23 juillet 2026", "Préparation et échaudage", "31"],
        ["Session 2", "19–20 août 2026", "Séchage et conditionnement", "29"],
        ["Session 3", "16–17 septembre 2026", "Tenue des registres et traçabilité", "26"],
    ])
    para(doc,
         "Le taux de satisfaction déclaré est de 91 %, mesuré par questionnaire en fin de "
         "session. Une quatrième session était prévue en septembre ; elle a été reportée "
         "faute de formateur disponible.")

    # ---------- 5. Difficultés ----------
    titre(doc, "5. Difficultés rencontrées et risques", 1)
    titre(doc, "5.1 Logistique d'export", 2)
    para(doc,
         "Les départs maritimes depuis le port de desserte ont été réduits à deux rotations "
         "mensuelles au lieu de quatre. Deux expéditions ont été décalées, pour un retard "
         "cumulé de 23 jours. Le surcoût de stockage s'élève à 6 800 € sur le trimestre.")

    titre(doc, "5.2 Trésorerie des producteurs", 2)
    para(doc,
         "Le délai moyen de paiement de 21 jours reste supérieur à l'objectif de 15 jours. "
         "Plusieurs producteurs ont signalé des difficultés de trésorerie en fin de campagne. "
         "Le risque est une perte de partenaires au profit de collecteurs informels payant "
         "au comptant.")

    titre(doc, "5.3 Disponibilité des formateurs", 2)
    para(doc,
         "L'équipe ne compte qu'un formateur qualifié sur la partie traçabilité. Son "
         "indisponibilité en septembre a entraîné le report d'une session.")

    # ---------- 6. Actions ----------
    titre(doc, "6. Actions à mener", 1)
    tableau(doc, ["Action", "Responsable", "Échéance"], [
        ["Renégocier le contrat de transport maritime", "Non désigné à ce jour", "30 novembre 2026"],
        ["Mettre en place une avance de campagne pour les producteurs",
         "R. Andriamalala", "15 décembre 2026"],
        ["Former un second formateur traçabilité", "Service formation", "Non définie"],
        ["Reprogrammer la session de formation reportée", "Service formation", "28 novembre 2026"],
        ["Clore les trois audits biologiques en écart", "Cabinet d'audit partenaire",
         "31 décembre 2026"],
    ])

    # ---------- 7. Prochaines étapes ----------
    titre(doc, "7. Prochaines étapes", 1)
    para(doc,
         "Le quatrième trimestre 2026 sera consacré à la phase d'analyse des données de "
         "campagne et à la préparation du rapport de filière, livrable attendu le "
         "15 décembre 2026.")
    para(doc,
         "Le comité de pilotage du 14 novembre 2026 devra arbitrer sur deux points : le "
         "financement de l'avance de campagne et le recrutement d'un second formateur.")
    para(doc,
         "Une cartographie complète des producteurs partenaires sera jointe au rapport de "
         "filière.")

    titre(doc, "Annexe — Méthodologie", 1)
    para(doc,
         "Les volumes proviennent des registres de pesée des sept centres de collecte. "
         "Le chiffre d'affaires est issu de la comptabilité analytique arrêtée au "
         "30 septembre 2026. Les taux de certification sont calculés sur les surfaces "
         "déclarées, non sur le nombre d'exploitations.")

    SORTIE.parent.mkdir(parents=True, exist_ok=True)
    doc.save(SORTIE)
    return SORTIE


if __name__ == "__main__":
    chemin = construire()
    print(f"Rapport généré : {chemin}")
