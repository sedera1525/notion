#!/usr/bin/env python3
"""Génère un second rapport fictif, pour un client SANS projet dans Notion.

Raison d'être : éprouver le « cas C » de la règle de rattachement du skill.
Quand le client du rapport n'a aucun projet dans « Projets / études », le skill
doit créer une page ; et si le même rapport revient corrigé, il doit retrouver
cette page et y ajouter une mise à jour, au lieu d'en créer une seconde.

Le client retenu est « Caldeira Solaire » (Énergie, La Réunion) : sa fiche
existe dans « Clients & prospects ». Le rapport est volontairement plus court
que celui de la démonstration principale : c'est une pièce de test, pas une
vitrine.

⚠️ Le cas C ne se rejoue pas tel quel
-------------------------------------
Au premier passage, la relation « Projets / études » de Caldeira Solaire était
vide — c'est ce qui déclenchait le cas C. Le skill a depuis créé la page
« Synthèse T3 2026 — Caldeira Solaire », qui est reliée à la fiche client. Elle
est donc devenue le projet de ce client, et un nouveau passage emprunte le
cas A, pas le cas C.

Pour éprouver à nouveau le cas C : supprimer d'abord cette page dans Notion.

⚠️ Le fichier produit n'est pas stable octet par octet
------------------------------------------------------
python-docx horodate les entrées de l'archive zip. Deux générations successives
donnent donc des empreintes SHA-256 différentes. Comme lancer_synthese.sh suit
les rapports par empreinte, **régénérer ce fichier le fait passer pour un
rapport corrigé** : le prochain traitement ajoutera une section « mise à jour »
dans Notion. Ne le régénérer que si c'est voulu.

Toutes les données sont inventées. Aucune entreprise, aucun montant et aucune
personne ne correspond à la réalité.

Usage : .venv/bin/python scripts/generer_rapport_cas_c.py
Sortie : input/rapport-t3-caldeira.docx
"""

from pathlib import Path

from _docx_commun import para, preparer, tableau, titre
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

SORTIE = Path(__file__).resolve().parent.parent / "input" / "rapport-t3-caldeira.docx"


def construire():
    doc = preparer(Document())

    # ---------- En-tête ----------
    t = doc.add_heading("Rapport d'activité trimestriel", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st = para(doc, "Caldeira Solaire — Exploitation du parc photovoltaïque", gras=True)
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    tableau(doc, ["Rubrique", "Information"], [
        ["Client", "Caldeira Solaire"],
        ["Pays", "La Réunion"],
        ["Période couverte", "1er juillet 2026 – 30 septembre 2026 (T3 2026)"],
        ["Auteur du rapport", "M. Payet, responsable exploitation"],
        ["Date de remise", "10 octobre 2026"],
        ["Référence", "CS-2026-T3-004"],
    ])
    para(doc, "Document de démonstration. Données entièrement fictives.", italique=True)

    # ---------- 1. Contexte ----------
    titre(doc, "1. Contexte et objectifs", 1)
    para(doc,
         "Caldeira Solaire exploite trois centrales photovoltaïques au sol dans le sud de "
         "La Réunion. Le présent rapport rend compte de l'exploitation du troisième "
         "trimestre 2026 et accompagne le dossier de consultation en cours sur l'étude "
         "d'impact de la future centrale.")
    para(doc,
         "Il fait suite au rapport du deuxième trimestre, remis le 9 juillet 2026. "
         "Le trimestre a été marqué par la mise en service d'une extension et par une "
         "tension sur l'approvisionnement en onduleurs.")

    # ---------- 2. Chiffres clés ----------
    titre(doc, "2. Chiffres clés du trimestre", 1)
    para(doc,
         "Les indicateurs sont comparés au deuxième trimestre 2026 lorsque la donnée "
         "antérieure est disponible.")
    tableau(doc, ["Indicateur", "T3 2026", "T2 2026", "Évolution"], [
        ["Production photovoltaïque", "4 820 MWh", "4 310 MWh", "+11,8 %"],
        ["Chiffre d'affaires", "578 000 €", "512 000 €", "+12,9 %"],
        ["Puissance installée", "12,6 MWc", "11,2 MWc", "+1,4 MWc"],
        ["Taux de disponibilité des onduleurs", "96,4 %", "94,1 %", "+2,3 points"],
        ["Incidents d'exploitation", "7", "11", "−4 incidents"],
        ["Coût d'exploitation au MWh", "18,70 €", "—", "—"],
    ])
    para(doc,
         "Le coût d'exploitation au MWh est calculé pour la première fois ce trimestre. "
         "Aucune comparaison n'est possible avec les périodes précédentes.")
    para(doc,
         "L'objectif interne de disponibilité des onduleurs est fixé à 97 %. "
         "Il n'est pas atteint.")

    # ---------- 3. Réalisations ----------
    titre(doc, "3. Réalisations de la période", 1)
    titre(doc, "3.1 Mise en service de l'extension sud", 2)
    para(doc,
         "L'extension de la centrale sud a été mise en service le 12 août 2026. Elle "
         "ajoute 1,4 MWc de puissance installée, portant le total du parc à 12,6 MWc. "
         "La réception technique a été prononcée sans réserve.")

    titre(doc, "3.2 Maintenance prédictive", 2)
    para(doc,
         "Le dispositif de suivi à distance couvre désormais 100 % des onduleurs, contre "
         "68 % au trimestre précédent. Sur les 7 incidents du trimestre, 5 ont été "
         "détectés avant l'arrêt de production.")

    titre(doc, "3.3 Nettoyage des modules", 2)
    para(doc,
         "Deux campagnes de nettoyage ont été menées, en juillet et en septembre 2026. "
         "Le gain de production mesuré après la seconde campagne est de 3,1 %.")

    # ---------- 4. Difficultés ----------
    titre(doc, "4. Difficultés rencontrées et risques", 1)
    titre(doc, "4.1 Approvisionnement en onduleurs", 2)
    para(doc,
         "Le délai de livraison des onduleurs de remplacement est passé de 6 à 14 semaines. "
         "Deux onduleurs de la centrale est fonctionnent en mode dégradé dans l'attente "
         "d'une pièce. La perte de production associée est estimée à 42 MWh sur le trimestre.")

    titre(doc, "4.2 Foncier de la future centrale", 2)
    para(doc,
         "La promesse de bail sur la parcelle visée n'est pas signée. Le propriétaire "
         "demande une révision du loyer. Le montant proposé n'est pas précisé dans les "
         "échanges à ce stade.")

    titre(doc, "4.3 Saison cyclonique", 2)
    para(doc,
         "La saison cyclonique débute en novembre. Le plan de mise en sécurité du parc "
         "existe, mais il n'a pas été testé depuis 2024.")

    # ---------- 5. Actions ----------
    titre(doc, "5. Actions à mener", 1)
    tableau(doc, ["Action", "Responsable", "Échéance"], [
        ["Constituer un stock tampon de deux onduleurs", "M. Payet", "30 novembre 2026"],
        ["Renégocier la promesse de bail", "Non désigné à ce jour", "15 décembre 2026"],
        ["Tester le plan de mise en sécurité cyclonique", "Service exploitation",
         "31 octobre 2026"],
        ["Fiabiliser le calcul du coût d'exploitation au MWh", "Service exploitation",
         "Non définie"],
    ])

    # ---------- 6. Prochaines étapes ----------
    titre(doc, "6. Prochaines étapes", 1)
    para(doc,
         "Le quatrième trimestre 2026 sera consacré à la préparation de la saison "
         "cyclonique et à la remise de l'étude d'impact de la future centrale, attendue "
         "le 20 décembre 2026.")
    para(doc,
         "Un comité d'exploitation est prévu le 18 novembre 2026 pour arbitrer le "
         "financement du stock tampon d'onduleurs.")

    titre(doc, "Annexe — Méthodologie", 1)
    para(doc,
         "Les volumes de production proviennent des compteurs de production des trois "
         "centrales. Le chiffre d'affaires est issu de la comptabilité arrêtée au "
         "30 septembre 2026. Le taux de disponibilité est calculé sur les heures "
         "d'ensoleillement exploitables, non sur les heures calendaires.")

    SORTIE.parent.mkdir(parents=True, exist_ok=True)
    doc.save(SORTIE)
    return SORTIE


if __name__ == "__main__":
    chemin = construire()
    print(f"Rapport généré : {chemin}")
